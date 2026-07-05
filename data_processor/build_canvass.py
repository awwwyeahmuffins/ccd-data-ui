"""
Build per-precinct canvass-coverage aggregates from a canvassing workbook.

Reads a Collin County canvassing export (an .xlsx with one row per precinct
carrying a "Canvassed Dem Voters %" and a "Dem Voters" universe count) and
emits ONLY aggregate counts/shares per base precinct to
data/tx/collin/<set>/profile/canvass.csv. The campaign / chair dashboards can
read this to show how much of each precinct's Democratic universe has already
been door-knocked. When the source workbook is absent (every machine but the
maintainer's), this prints a note and exits 0 without writing anything — the
frontend renders N/A.

The source workbook is operational field data, not part of the committed data
pipeline: drop it at the repo root as "canvass_source.xlsx" (gitignored, like
the voter file) or pass --source. The OUTPUT (canvass.csv) is aggregate-only
and safe to commit and deploy.

CAVEAT — boundary vintage: the reference export is keyed to 2024 precinct
numbers. Precinct numbering is largely stable across the 2024/2026 boundary
sets but is not guaranteed identical; --set controls which profile dir the
CSV lands in (default 2026, where the dashboards look).

Usage:
    python3 data_processor/build_canvass.py [--source PATH] [--set 2026]
                                            [--sheet NAME] [--dry-run]
"""

import argparse
import csv
import sys
from datetime import date
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
DEFAULT_SOURCE = REPO / "canvass_source.xlsx"

# Source column headers (exact, row 1 of the sheet).
COL_PRECINCT = "Precinct"
COL_CANVASSED_PCT = "Canvassed Dem Voters %"
COL_DEM_VOTERS = "Dem Voters"


def base_precinct(raw):
    """Sheet precinct value (1.0, "23") -> 23; None if unparseable/blank."""
    if raw is None:
        return None
    try:
        return int(float(str(raw).split(".")[0]) if "." not in str(raw)
                   else float(raw))
    except (ValueError, TypeError):
        return None


def load_rows(source, sheet_name):
    try:
        import openpyxl
    except ImportError:
        sys.exit("ERROR: openpyxl is required to read the workbook "
                 "(pip install openpyxl).")
    wb = openpyxl.load_workbook(source, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb.worksheets[0]

    header = {}
    for c in range(1, ws.max_column + 1):
        name = ws.cell(1, c).value
        if isinstance(name, str):
            header[name.strip()] = c
    for needed in (COL_PRECINCT, COL_CANVASSED_PCT, COL_DEM_VOTERS):
        if needed not in header:
            sys.exit(f"ERROR: column {needed!r} not found in sheet "
                     f"{ws.title!r}. Found: {sorted(header)}")

    rows = []
    for r in range(2, ws.max_row + 1):
        pct = base_precinct(ws.cell(r, header[COL_PRECINCT]).value)
        if pct is None:
            continue
        share = ws.cell(r, header[COL_CANVASSED_PCT]).value
        dem = ws.cell(r, header[COL_DEM_VOTERS]).value
        share = float(share) if isinstance(share, (int, float)) else None
        dem = int(dem) if isinstance(dem, (int, float)) else None
        canvassed = (round(share * dem) if share is not None and dem is not None
                     else "")
        rows.append({
            "precinct": pct,
            "dem_voters": dem if dem is not None else "",
            "canvassed_share": f"{share:.4f}" if share is not None else "",
            "canvassed_voters": canvassed,
            "generated_on": date.today().isoformat(),
        })
    rows.sort(key=lambda r: r["precinct"])
    return rows, ws.title


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[1])
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE,
                        help=f"canvassing .xlsx (default: {DEFAULT_SOURCE.name} "
                             "at repo root, gitignored)")
    parser.add_argument("--set", default="2026", dest="boundary_set",
                        help="boundary set to write under (default: 2026)")
    parser.add_argument("--sheet", default=None,
                        help="worksheet name (default: first sheet)")
    parser.add_argument("--dry-run", action="store_true",
                        help="parse and report, write nothing")
    args = parser.parse_args()

    if not args.source.exists():
        print(f"Canvass workbook not present ({args.source}); nothing to do. "
              "The dashboard renders N/A without canvass.csv.")
        return

    output_path = (REPO / "data" / "tx" / "collin" / args.boundary_set /
                   "profile" / "canvass.csv")
    if not output_path.parent.is_dir():
        sys.exit(f"ERROR: no such profile dir: {output_path.parent} "
                 f"(unknown boundary set '{args.boundary_set}'?)")

    rows, sheet_title = load_rows(args.source, args.sheet)
    covered = [r for r in rows if r["canvassed_share"] != ""]
    print(f"Parsed sheet {sheet_title!r}: {len(rows)} precincts, "
          f"{len(covered)} with a canvass share.")
    if covered:
        total_dem = sum(r["dem_voters"] for r in covered
                        if r["dem_voters"] != "")
        total_canvassed = sum(r["canvassed_voters"] for r in covered
                              if r["canvassed_voters"] != "")
        if total_dem:
            print(f"Dem universe: {total_dem:,} | canvassed: "
                  f"{total_canvassed:,} ({total_canvassed / total_dem:.1%})")
        least = sorted(covered, key=lambda r: float(r["canvassed_share"]))[:5]
        print("Least-canvassed precincts:")
        for r in least:
            print(f"  precinct {r['precinct']:>4}: "
                  f"{float(r['canvassed_share']):.1%} of "
                  f"{r['dem_voters']} Dem voters")

    if args.dry_run:
        print("\n--dry-run: nothing written.")
        return

    with open(output_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=[
            "precinct", "dem_voters", "canvassed_share",
            "canvassed_voters", "generated_on",
        ])
        writer.writeheader()
        writer.writerows(rows)
    print(f"\nWrote {output_path.relative_to(REPO)} ({len(rows)} rows, "
          "aggregate counts only).")


if __name__ == "__main__":
    main()
